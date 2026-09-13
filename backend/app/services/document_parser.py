import csv
from pathlib import Path

from app.rag.chunking import ParsedPage


def parse_document(path: str, mime_type: str) -> list[ParsedPage]:
    source = Path(path)
    suffix = source.suffix.lower()
    if suffix == ".pdf" or mime_type == "application/pdf":
        from pypdf import PdfReader

        return [
            ParsedPage(page_number=i, text=page.extract_text() or "")
            for i, page in enumerate(PdfReader(source).pages, start=1)
        ]
    if suffix == ".docx":
        from docx import Document

        doc = Document(source)
        blocks = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            blocks.extend(
                " | ".join(cell.text.strip() for cell in row.cells)
                for row in table.rows
            )
        text = "\n".join(block for block in blocks if block.strip())
        return [ParsedPage(page_number=1, text=text)]
    if suffix == ".csv":
        with source.open(encoding="utf-8-sig", newline="") as handle:
            rows = list(csv.reader(handle))
        return [ParsedPage(page_number=1, text="\n".join(" | ".join(row) for row in rows))]
    if suffix == ".txt" or mime_type.startswith("text/"):
        return [ParsedPage(page_number=1, text=source.read_text(encoding="utf-8-sig"))]
    raise ValueError(f"Unsupported document type: {suffix or mime_type}")
